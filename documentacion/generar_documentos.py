from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DOCS_DIR = Path(__file__).resolve().parent


def _sombrear(celda, color):
    tc_pr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def _configurar_estilos(doc):
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(10.5)
    estilo.paragraph_format.space_after = Pt(4)
    for seccion in doc.sections:
        seccion.top_margin = Cm(2.2)
        seccion.bottom_margin = Cm(2.2)
        seccion.left_margin = Cm(2.2)
        seccion.right_margin = Cm(2.2)


def _pie(doc):
    seccion = doc.sections[0]
    parrafo = seccion.footer.paragraphs[0]
    parrafo.text = "AutoPark - Documentación de defensa"
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _portada(doc, numero, titulo, subtitulo):
    parrafo = doc.add_paragraph()
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = parrafo.add_run("AutoPark")
    run.bold = True
    run.font.size = Pt(30)
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    parrafo = doc.add_paragraph()
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = parrafo.add_run("Sistema de Gestión de Estacionamiento")
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x40, 0x40, 0x40)

    parrafo = doc.add_heading(titulo, level=0)
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    parrafo = doc.add_paragraph(subtitulo)
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    parrafo.runs[0].italic = True

    _t(
        doc,
        ["Dato", "Detalle"],
        [
            ["Materia", "Programación I / II"],
            ["Proyecto", "2.º Examen Parcial - AutoPark"],
            ["Documento", numero],
            ["Modalidad", "Dupla - Desafío Nivel 2: GUI con tkinter"],
            ["Fecha", date.today().strftime("%d/%m/%Y")],
        ],
    )
    doc.add_page_break()


def _nuevo(numero, titulo, subtitulo):
    doc = Document()
    _configurar_estilos(doc)
    _pie(doc)
    _portada(doc, numero, titulo, subtitulo)
    doc.core_properties.title = titulo
    doc.core_properties.subject = "AutoPark - Documentación de defensa"
    doc.core_properties.author = "Equipo AutoPark"
    return doc


def _h(doc, texto, nivel=1):
    return doc.add_heading(texto, level=nivel)


def _p(doc, texto, negrita=False, cursiva=False, tamano=None):
    parrafo = doc.add_paragraph()
    run = parrafo.add_run(texto)
    run.bold = negrita
    run.italic = cursiva
    if tamano is not None:
        run.font.size = Pt(tamano)
    return parrafo


def _b(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def _n(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


def _t(doc, encabezados, filas):
    tabla = doc.add_table(rows=1, cols=len(encabezados))
    tabla.style = "Table Grid"
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabla.autofit = True
    celdas_encabezado = tabla.rows[0].cells
    for indice, texto in enumerate(encabezados):
        celdas_encabezado[indice].text = ""
        run = celdas_encabezado[indice].paragraphs[0].add_run(texto)
        run.bold = True
        _sombrear(celdas_encabezado[indice], "D9E2F3")
    for fila in filas:
        celdas = tabla.add_row().cells
        for indice, valor in enumerate(fila):
            celdas[indice].text = str(valor)
    return tabla


def _c(doc, texto):
    parrafo = doc.add_paragraph()
    parrafo.paragraph_format.space_after = Pt(6)
    p_pr = parrafo._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F2F2")
    p_pr.append(shd)
    for indice, linea in enumerate(texto.strip("\n").split("\n")):
        if indice:
            parrafo.add_run().add_break()
        run = parrafo.add_run(linea)
        run.font.name = "Consolas"
        run.font.size = Pt(9)


def _qa(doc, pares):
    for pregunta, respuesta in pares:
        parrafo = doc.add_paragraph()
        run = parrafo.add_run("Pregunta: ")
        run.bold = True
        parrafo.add_run(pregunta)
        parrafo = doc.add_paragraph()
        run = parrafo.add_run("Respuesta: ")
        run.bold = True
        parrafo.add_run(respuesta)


def _check(doc, items):
    for item in items:
        doc.add_paragraph("[ ] " + item, style="List Bullet")


def _nota(doc, texto):
    parrafo = doc.add_paragraph()
    run = parrafo.add_run(texto)
    run.italic = True


def _guardar(doc, nombre):
    ruta = DOCS_DIR / nombre
    doc.save(ruta)
    print(f"Generado: {ruta.name}")


def documento_00():
    doc = _nuevo(
        "Guía general",
        "Guía general y guion de defensa",
        "Vista completa del sistema, arquitectura y recorrido de la defensa",
    )

    _h(doc, "1. Qué vamos a defender")
    _p(
        doc,
        "AutoPark es una aplicación modular de consola con una interfaz gráfica "
        "para el cobro manual. El sistema base permite registrar un cobro manual, "
        "procesar un lote de movimientos desde JSON y exportar los resultados a CSV. "
        "El desafío de nivel 2 elegido es la GUI con tkinter.",
    )
    _b(
        doc,
        [
            "Sistema base obligatorio: cobro manual, lote, validaciones, recursión y CSV.",
            "Desafío de nivel 2: ventana gráfica para el Cobro Manual en Ventanilla.",
            "No se usa SQLite: la persistencia obligatoria se resuelve con CSV.",
            "La GUI no duplica la lógica: reutiliza los módulos de cálculo y CSV.",
        ],
    )

    _h(doc, "2. Arquitectura en una vista")
    _t(
        doc,
        ["Archivo", "Responsabilidad"],
        [
            ["main.py", "Menú principal, submenú manual y coordinación general"],
            ["config.py", "Rutas, separador, regex y constantes"],
            ["utilidades.py", "Formato de montos, fecha e ID manual"],
            ["json_repo.py", "Carga y validación de tarifas y movimientos"],
            ["validaciones.py", "Patentes, tipos, horas, días y lavado"],
            ["calculo.py", "Recursión, topes y liquidación"],
            ["lote.py", "Procesamiento y separación de válidos/rechazados"],
            ["csv_repo.py", "Escritura de reportes y acumulación manual"],
            ["manual_consola.py", "Cobro manual y consulta por consola"],
            ["gui_manual.py", "Ventana tkinter, eventos y vista previa"],
            ["pruebas.py", "Pruebas automáticas con assert"],
        ],
    )

    _h(doc, "3. Flujo del cobro manual")
    _n(
        doc,
        [
            "El usuario elige la opción 1 del menú principal.",
            "Elige ventana gráfica o consola.",
            "Ingresa patente, tipo, días, horarios y lavado.",
            "El sistema normaliza y valida los datos.",
            "Calcula base, bonificación, lavado y total.",
            "Muestra el desglose y guarda en liquidacion-manual.csv.",
            "Limpia el formulario o vuelve al menú.",
        ],
    )

    _h(doc, "4. Flujo del procesamiento por lote")
    _n(
        doc,
        [
            "Se cargan tarifas.json y movimientos.json.",
            "Se recorre cada movimiento en orden.",
            "Se validan ID, patente, tipo, horario y lavado.",
            "Los inválidos se derivan a rechazados con su motivo.",
            "Los válidos se liquidan con la misma función del cobro manual.",
            "Se escriben liquidacion.csv y rechazados.csv.",
            "Se muestra el resumen: procesados, rechazados y total recaudado.",
        ],
    )

    _h(doc, "5. Orden de lectura de los documentos")
    _t(
        doc,
        ["Documento", "Tema", "Para qué sirve"],
        [
            ["00", "Guía general", "Visión completa y guion de defensa"],
            ["01", "Preparación y JSON", "Explicar estructura, archivos y carga"],
            ["02", "Validaciones", "Defender patentes, horas y tipos"],
            ["03", "Cálculo y recursión", "Defender la liquidación y el acumulador"],
            ["04", "CSV", "Defender persistencia y acumulación"],
            ["05", "Lote", "Defender los 30 válidos y 3 rechazados"],
            ["06", "Consola", "Defender el flujo de respaldo"],
            ["07", "GUI tkinter", "Defender eventos, widgets y total en vivo"],
            ["08", "Menú e integración", "Defender la unión de todos los módulos"],
            ["09", "Pruebas y defensa", "Repasar casos, demo y preguntas"],
        ],
    )

    _h(doc, "6. Guion de defensa de 5 minutos")
    _n(
        doc,
        [
            "Presentar la arquitectura: main, validaciones, cálculo, CSV, lote y GUI.",
            "Ejecutar la opción 2 y mostrar 30 procesados, 3 rechazados y $129.000.",
            "Abrir la GUI y cargar auto con 210 minutos y lavado: total $7.400.",
            "Cargar camioneta con 2 días y lavado: total $55.500.",
            "Probar una patente inválida y mostrar el mensaje de error.",
            "Mostrar la fila guardada en liquidacion-manual.csv.",
            "Explicar que la GUI solo captura datos y llama a los módulos existentes.",
        ],
    )

    _h(doc, "7. Conceptos clave")
    _b(
        doc,
        [
            "Modularidad: cada archivo tiene una responsabilidad concreta.",
            "JSON: configuración y lote de datos de entrada.",
            "CSV: persistencia simple, legible y compatible con planillas.",
            "Recursión: cálculo de bloques con acumulador de minutos tarifados.",
            "Eventos: tkinter reacciona a cambios de variables y clics.",
            "Validación: normalizar primero, validar después y recién entonces calcular.",
        ],
    )

    _h(doc, "8. Preguntas generales")
    _qa(
        doc,
        [
            (
                "¿Qué hace el sistema?",
                "Registra cobros manuales, procesa un lote desde JSON y genera reportes CSV. "
                "Además, ofrece una GUI para el cobro manual.",
            ),
            (
                "¿Por qué está dividido en módulos?",
                "Para separar responsabilidades, facilitar las pruebas y evitar duplicar "
                "la lógica entre consola, GUI y lote.",
            ),
            (
                "¿Qué desafío de nivel 2 eligieron?",
                "La GUI con tkinter para el módulo de Cobro Manual en Ventanilla.",
            ),
            (
                "¿Por qué no usaron SQLite?",
                "Porque el desafío elegido es GUI y la consigna base solo exige CSV. "
                "SQLite habría sumado tiempo y no era necesario.",
            ),
            (
                "¿Cómo se guardan los datos?",
                "El lote se escribe en liquidacion.csv y rechazados.csv. "
                "El cobro manual se acumula en liquidacion-manual.csv.",
            ),
            (
                "¿Qué pasa si un JSON está mal?",
                "json_repo valida estructura y claves. main muestra el error y no procesa.",
            ),
            (
                "¿Cómo probaron el sistema?",
                "Con pruebas automáticas en pruebas.py y una prueba manual de la GUI.",
            ),
        ],
    )

    _h(doc, "9. Checklist general")
    _check(
        doc,
        [
            "Tener abierta la carpeta autopark.",
            "Tener a mano tarifas.json y movimientos.json.",
            "Ejecutar python pruebas.py antes de la defensa.",
            "Probar la opción 2 y verificar 30/3/$129.000.",
            "Probar la GUI con 210 minutos y lavado.",
            "Tener liquidacion-manual.csv visible.",
            "Saber explicar la recursión con el ejemplo de 135 minutos.",
        ],
    )

    _h(doc, "10. Nota sobre el enunciado")
    _nota(
        doc,
        "El ENUNCIADO.pdf dice Modalidad Individual. El catálogo "
        "DESAFIOS_VALOR_AGREGADO.pdf indica que la modalidad en dupla exige un "
        "desafío de NIVEL 2. Para la defensa, llevar impreso el catálogo y aclarar "
        "que se optó por la GUI, que es el desafío 2 del nivel 2.",
    )
    _nota(
        doc,
        "El DOCX del profesor presenta una recursión que no coincide con el ejemplo "
        "explícito de 135 minutos del enunciado. El sistema sigue el enunciado: "
        "135 minutos = 4 bloques de 30 + 1 hora adicional = $4.000.",
    )

    _guardar(doc, "00_indice_general.docx")


def documento_01():
    doc = _nuevo(
        "01 de 09",
        "Preparación, JSON y configuración",
        "Estructura del proyecto, archivos de entrada y carga de datos",
    )

    _h(doc, "1. Objetivo de la fase")
    _p(
        doc,
        "Crear la base del proyecto, definir las rutas y cargar los datos de entrada "
        "desde JSON. Esta fase explica por qué el sistema no tiene tarifas ni "
        "movimientos hardcodeados.",
    )

    _h(doc, "2. Archivos que intervienen")
    _t(
        doc,
        ["Archivo", "Función"],
        [
            ["config.py", "Rutas absolutas, separador, regex y constantes"],
            ["utilidades.py", "Formato de montos, fecha e ID"],
            ["json_repo.py", "Carga y validación de los JSON"],
            ["tarifas.json", "Configuración y precios"],
            ["movimientos.json", "Lote de 33 movimientos"],
        ],
    )

    _h(doc, "3. Estructura del proyecto")
    _c(
        doc,
        r"""
autopark/
├── main.py
├── config.py
├── utilidades.py
├── json_repo.py
├── validaciones.py
├── calculo.py
├── lote.py
├── csv_repo.py
├── manual_consola.py
├── gui_manual.py
├── pruebas.py
├── README.md
├── tarifas.json
└── movimientos.json
""",
    )
    _p(
        doc,
        "Los archivos generados son liquidacion.csv, rechazados.csv y "
        "liquidacion-manual.csv. No se usa carpeta data ni base de datos.",
    )

    _h(doc, "4. tarifas.json")
    _p(
        doc,
        "El archivo tiene dos bloques: configuracion y tarifas. La configuración "
        "define los minutos de cada tramo. Las tarifas definen precio por bloque, "
        "precio por hora adicional, tope diario y precio de lavado por categoría.",
    )
    _c(
        doc,
        r"""
{
  "configuracion": {
    "minutos_bloque_inicial": 30,
    "minutos_limite_primer_tramo": 120,
    "minutos_hora_adicional": 60,
    "minutos_dia_completo": 1440
  },
  "tarifas": {
    "moto": {
      "precio_bloque_30min": 400,
      "precio_hora_adicional": 700,
      "tope_dia": 9000,
      "precio_lavado": 2500
    },
    "auto": {
      "precio_bloque_30min": 700,
      "precio_hora_adicional": 1200,
      "tope_dia": 16000,
      "precio_lavado": 5000
    },
    "camioneta": {
      "precio_bloque_30min": 1000,
      "precio_hora_adicional": 1800,
      "tope_dia": 24000,
      "precio_lavado": 7500
    }
  }
}
""",
    )

    _h(doc, "5. movimientos.json")
    _p(
        doc,
        "El lote tiene fecha_operacion y una lista de 33 movimientos. Cada movimiento "
        "tiene id, patente, tipo_vehiculo, entrada, salida y solicito_lavado.",
    )
    _c(
        doc,
        r"""
{
  "fecha_operacion": "2026-10-19",
  "movimientos": [
    {
      "id": "MOV001",
      "patente": "AA123BB",
      "tipo_vehiculo": "auto",
      "entrada": "08:00",
      "salida": "08:25",
      "solicito_lavado": false
    }
  ]
}
""",
    )

    _h(doc, "6. config.py")
    _p(
        doc,
        "Se usa Path(__file__).resolve().parent para que las rutas funcionen aunque "
        "el programa se ejecute desde otra carpeta. El separador CSV queda en una "
        "sola constante para poder cambiarlo sin tocar todos los módulos.",
    )
    _c(
        doc,
        r"""
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent

ARCHIVO_TARIFAS = BASE_DIR / "tarifas.json"
ARCHIVO_MOVIMIENTOS = BASE_DIR / "movimientos.json"
CSV_LIQUIDACION = BASE_DIR / "liquidacion.csv"
CSV_RECHAZADOS = BASE_DIR / "rechazados.csv"
CSV_MANUAL = BASE_DIR / "liquidacion-manual.csv"

SEPARADOR = ","
FORMATO_FECHA = "%Y-%m-%d"
PATRON_PATENTE = r"^(?:[A-Z]{3}\d{3}|[A-Z]{2}\d{3}[A-Z]{2})$"
TITULO_VENTANA = "AutoPark - Cobro Manual"
""",
    )

    _h(doc, "7. json_repo.py")
    _p(
        doc,
        "json_repo no solo lee el archivo: también comprueba que existan las claves "
        "obligatorias y que los valores numéricos sean válidos. Así el resto del "
        "sistema trabaja con datos confiables.",
    )
    _c(
        doc,
        r"""
def cargar_json(ruta):
    with open(ruta, "r", encoding="utf-8") as archivo:
        return json.load(archivo)


def cargar_tarifas(ruta=ARCHIVO_TARIFAS):
    datos = cargar_json(ruta)
    if "configuracion" not in datos or "tarifas" not in datos:
        raise ValueError("El archivo de tarifas debe tener 'configuracion' y 'tarifas'.")
    configuracion = datos["configuracion"]
    tarifas = datos["tarifas"]
    for clave in CLAVES_CONFIGURACION:
        valor = configuracion.get(clave)
        if isinstance(valor, bool) or not isinstance(valor, int) or valor <= 0:
            raise ValueError(f"La configuración '{clave}' debe ser un entero positivo.")
    return configuracion, tarifas
""",
    )

    _h(doc, "8. utilidades.py")
    _p(
        doc,
        "formato_pesos se usa solo para mostrar en pantalla. En los CSV se guardan "
        "números con dos decimales para que puedan sumarse o abrirse en una planilla.",
    )
    _c(
        doc,
        r"""
def formato_pesos(valor):
    return f"$ {float(valor):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def generar_id_manual():
    return "MAN-" + datetime.now().strftime("%Y%m%d%H%M%S%f")
""",
    )

    _h(doc, "9. Cómo defender esta fase")
    _b(
        doc,
        [
            "El sistema no tiene precios hardcodeados: todo sale de tarifas.json.",
            "Las rutas son absolutas a partir del archivo config.py.",
            "json_repo valida estructura y tipos antes de calcular.",
            "Los JSON están en la raíz del proyecto para que sean fáciles de encontrar.",
            "Si el docente cambia un precio, no hay que tocar el código.",
        ],
    )

    _h(doc, "10. Preguntas del docente")
    _qa(
        doc,
        [
            (
                "¿Por qué usar JSON?",
                "Porque separa la configuración y los datos del código, y porque es "
                "el formato pedido por la consigna.",
            ),
            (
                "¿Qué pasa si falta una clave?",
                "json_repo lanza ValueError con un mensaje claro y main no continúa.",
            ),
            (
                "¿Por qué Path(__file__)?",
                "Para que las rutas se resuelvan desde la carpeta del proyecto y no "
                "dependan de la carpeta desde donde se ejecuta Python.",
            ),
            (
                "¿Qué pasa si fecha_operacion es inválida?",
                "Se valida con datetime.strptime y se informa el error.",
            ),
            (
                "¿Los CSV usan coma o punto y coma?",
                "Coma por defecto. Está centralizado en config.py por si se necesita cambiar.",
            ),
        ],
    )

    _h(doc, "11. Checklist de la fase")
    _check(
        doc,
        [
            "Existen tarifas.json y movimientos.json.",
            "El JSON tiene 33 movimientos.",
            "config.py usa BASE_DIR.",
            "json_repo valida claves y tipos.",
            "main.py informa si no puede cargar tarifas.json.",
        ],
    )

    _guardar(doc, "01_preparacion_json_configuracion.docx")


def documento_02():
    doc = _nuevo(
        "02 de 09",
        "Validaciones",
        "Normalización de patentes, tipos, horarios, días y lavado",
    )

    _h(doc, "1. Objetivo de la fase")
    _p(
        doc,
        "Recibir datos escritos por una persona o leídos de un JSON y convertirlos a "
        "un formato confiable. La regla es siempre la misma: primero normalizar, "
        "después validar y recién al final calcular.",
    )

    _h(doc, "2. Patentes")
    _p(
        doc,
        "Se eliminan todos los espacios y se pasa todo a mayúsculas. Luego se valida "
        "con una expresión regular que acepta solo los dos formatos oficiales.",
    )
    _t(
        doc,
        ["Entrada", "Normalizada", "Válida"],
        [
            ["AA123BB", "AA123BB", "Sí"],
            ["abc123", "ABC123", "Sí"],
            [" BF 901 GG ", "BF901GG", "Sí"],
            ["PATENTE99", "PATENTE99", "No"],
            ["AB1234", "AB1234", "No"],
        ],
    )
    _c(
        doc,
        r"""
def normalizar_patente(texto):
    if not isinstance(texto, str):
        return ""
    return re.sub(r"\s+", "", texto).upper()


def patente_valida(patente):
    if not isinstance(patente, str):
        return False
    return re.fullmatch(PATRON_PATENTE, patente) is not None
""",
    )

    _h(doc, "3. Tipos de vehículo")
    _p(
        doc,
        "El tipo se normaliza con strip y lower. Después se comprueba que exista "
        "dentro de las tarifas cargadas. Esto permite que el sistema acepte "
        "' MOTO ', 'AUTO' o ' camioneta' sin hardcodear cada variante.",
    )
    _c(
        doc,
        r"""
def normalizar_tipo(texto, tarifas):
    if not isinstance(texto, str):
        return None
    tipo = texto.strip().lower()
    return tipo if tipo in tarifas else None
""",
    )

    _h(doc, "4. Horarios")
    _p(
        doc,
        "Se acepta H:MM y HH:MM, entre 00:00 y 23:59. La función devuelve minutos "
        "desde las 00:00 para que el cálculo sea una simple resta.",
    )
    _t(
        doc,
        ["Entrada", "Resultado"],
        [
            ["08:00", "480"],
            ["8:00", "480"],
            ["23:59", "1439"],
            ["24:00", "None"],
            ["08:60", "None"],
        ],
    )
    _c(
        doc,
        r"""
def parsear_hora(texto):
    if not isinstance(texto, str):
        return None
    texto = texto.strip()
    if re.fullmatch(r"(?:[01]?\d|2[0-3]):[0-5]\d", texto) is None:
        return None
    hora, minuto = texto.split(":")
    return int(hora) * 60 + int(minuto)


def minutos_entre(entrada, salida):
    minutos_entrada = parsear_hora(entrada)
    minutos_salida = parsear_hora(salida)
    if minutos_entrada is None or minutos_salida is None:
        raise ValueError("Hora inválida. Usá HH:MM.")
    if minutos_salida <= minutos_entrada:
        raise ValueError("La salida debe ser posterior a la entrada.")
    return minutos_salida - minutos_entrada
""",
    )

    _h(doc, "5. Días y lavado")
    _p(
        doc,
        "Los días deben ser un entero mayor o igual a 0. El lavado se interpreta "
        "como booleano; en consola se pregunta S/N y en la GUI se usa un checkbox.",
    )
    _c(
        doc,
        r"""
def parsear_dias(texto):
    if isinstance(texto, bool):
        raise ValueError("Días debe ser un entero mayor o igual a 0.")
    if isinstance(texto, int):
        dias = texto
    elif isinstance(texto, str):
        texto = texto.strip()
        if not texto.isdigit():
            raise ValueError("Días debe ser un entero mayor o igual a 0.")
        dias = int(texto)
    else:
        raise ValueError("Días debe ser un entero mayor o igual a 0.")
    if dias < 0:
        raise ValueError("Días debe ser un entero mayor o igual a 0.")
    return dias


def parsear_lavado(texto):
    if isinstance(texto, bool):
        return texto
    if isinstance(texto, str):
        respuesta = texto.strip().lower()
        if respuesta in ("s", "si", "sí", "true", "1"):
            return True
        if respuesta in ("n", "no", "false", "0", ""):
            return False
    raise ValueError("Respuesta de lavado inválida. Usá S o N.")
""",
    )

    _h(doc, "6. Mensajes de error")
    _t(
        doc,
        ["Situación", "Mensaje"],
        [
            ["Patente mal escrita", "Patente inválida. Usá ABC123 o AA123BB."],
            ["Tipo inexistente", "Tipo de vehículo inválido."],
            ["Hora mal escrita", "Hora inválida. Usá HH:MM."],
            ["Salida anterior o igual", "La salida debe ser posterior a la entrada."],
            ["Días negativos o decimales", "Días debe ser un entero mayor o igual a 0."],
            ["Lavado mal respondido", "Respuesta de lavado inválida. Usá S o N."],
        ],
    )

    _h(doc, "7. Cómo defender esta fase")
    _b(
        doc,
        [
            "La normalización evita rechazar datos válidos solo por espacios o mayúsculas.",
            "La regex de patente es estricta: no acepta letras ni números de más.",
            "Las horas se convierten a minutos para simplificar la resta.",
            "El lote y la GUI usan las mismas funciones de validación.",
            "Los errores se muestran con mensajes claros, no con excepciones sin contexto.",
        ],
    )

    _h(doc, "8. Preguntas del docente")
    _qa(
        doc,
        [
            (
                "¿Por qué normalizar antes de validar?",
                "Para que ' bf 901 gg ' sea válida sin obligar al usuario a escribir "
                "exactamente en mayúsculas y sin espacios.",
            ),
            (
                "¿Qué patentes rechaza?",
                "Cualquiera que no sea 3 letras + 3 números o 2 letras + 3 números + 2 letras.",
            ),
            (
                "¿Qué pasa si la salida es igual a la entrada?",
                "Se rechaza porque no hay tiempo de estadía positivo.",
            ),
            (
                "¿Se valida el lavado en el lote?",
                "Sí. Si viene como texto se interpreta; si es inválido, el movimiento se rechaza.",
            ),
            (
                "¿Qué pasa si el tipo viene con espacios?",
                "Se normaliza con strip y lower, por eso ' MOTO ' se acepta como moto.",
            ),
        ],
    )

    _h(doc, "9. Checklist de la fase")
    _check(
        doc,
        [
            "Probar una patente tradicional y una Mercosur.",
            "Probar una patente con espacios.",
            "Probar una hora con un dígito y con dos.",
            "Probar una salida anterior a la entrada.",
            "Probar un tipo con mayúsculas y espacios.",
        ],
    )

    _guardar(doc, "02_validaciones.docx")


def documento_03():
    doc = _nuevo(
        "03 de 09",
        "Cálculo y recursión",
        "Bloques de 30 minutos, horas adicionales, topes y lavado",
    )

    _h(doc, "1. Objetivo de la fase")
    _p(
        doc,
        "Calcular el importe de una estadía respetando los tramos, el tope diario y "
        "las reglas de lavado. La recursión se usa para recorrer los bloques de tiempo "
        "sin escribir un bucle manual.",
    )

    _h(doc, "2. Regla de negocio")
    _b(
        doc,
        [
            "Primeros 120 minutos: bloques de 30 minutos.",
            "Desde el minuto 121: horas completas o fracción.",
            "El estacionamiento no puede superar el tope diario.",
            "Días mayores a 0: días * tope_dia, sin horarios.",
            "Lavado hasta 120 minutos: estacionamiento bonificado al 100%.",
            "Lavado con más de 120 minutos: se cobra solo el excedente.",
        ],
    )

    _h(doc, "3. Recursión con acumulador")
    _p(
        doc,
        "La función recibe los minutos que faltan cobrar y cuántos minutos ya fueron "
        "tarifados. Mientras no llegue a 120, cobra bloques de 30. Después, cobra "
        "horas de 60. El caso base es cuando ya no quedan minutos.",
    )
    _c(
        doc,
        r"""
def calcular_estacionamiento_recursivo(
    minutos_restantes,
    tipo,
    tarifas,
    configuracion,
    minutos_tarifados=0,
):
    if tipo not in tarifas:
        raise ValueError("Tipo de vehículo inválido.")
    if minutos_restantes <= 0:
        return 0

    tarifa = tarifas[tipo]
    bloque = configuracion["minutos_bloque_inicial"]
    limite = configuracion["minutos_limite_primer_tramo"]
    hora = configuracion["minutos_hora_adicional"]

    if minutos_tarifados < limite:
        return tarifa["precio_bloque_30min"] + calcular_estacionamiento_recursivo(
            minutos_restantes - bloque,
            tipo,
            tarifas,
            configuracion,
            minutos_tarifados + bloque,
        )

    return tarifa["precio_hora_adicional"] + calcular_estacionamiento_recursivo(
        minutos_restantes - hora,
        tipo,
        tarifas,
        configuracion,
        minutos_tarifados + hora,
    )
""",
    )

    _h(doc, "4. Ejemplo paso a paso: 135 minutos de auto")
    _t(
        doc,
        ["Minutos restantes", "Minutos tarifados", "Tramo", "Importe"],
        [
            ["135", "0", "Bloque 30", "700"],
            ["105", "30", "Bloque 30", "700"],
            ["75", "60", "Bloque 30", "700"],
            ["45", "90", "Bloque 30", "700"],
            ["15", "120", "Hora adicional", "1200"],
            ["-45", "180", "Caso base", "0"],
        ],
    )
    _p(doc, "Total: 4 bloques de 30 + 1 hora adicional = $4.000.")
    _nota(
        doc,
        "Este resultado coincide con el ejemplo explícito del ENUNCIADO.pdf. "
        "Si el docente menciona otra recursión del DOCX, se explica que se priorizó "
        "el ejemplo numérico del enunciado.",
    )

    _h(doc, "5. Tope diario")
    _p(
        doc,
        "La recursión calcula el costo por tiempo. Luego calcular_estacionamiento "
        "aplica el tope de la categoría con min(). Para un auto, el tope es $16.000.",
    )
    _c(
        doc,
        r"""
def calcular_estacionamiento(minutos, tipo, tarifas, configuracion):
    if tipo not in tarifas:
        raise ValueError("Tipo de vehículo inválido.")
    if minutos <= 0:
        return 0
    total = calcular_estacionamiento_recursivo(minutos, tipo, tarifas, configuracion)
    return min(total, tarifas[tipo]["tope_dia"])
""",
    )

    _h(doc, "6. Lavado y excedente")
    _p(
        doc,
        "Si hay lavado y la estadía es de hasta 120 minutos, el estacionamiento se "
        "bonifica completo. Si supera 120, se calcula solo el excedente usando la "
        "recursión con minutos_tarifados=120.",
    )
    _c(
        doc,
        r"""
def calcular_excedente(minutos, tipo, tarifas, configuracion):
    limite = configuracion["minutos_limite_primer_tramo"]
    if minutos <= limite:
        return 0
    total = calcular_estacionamiento_recursivo(
        minutos - limite,
        tipo,
        tarifas,
        configuracion,
        minutos_tarifados=limite,
    )
    return min(total, tarifas[tipo]["tope_dia"])
""",
    )

    _h(doc, "7. Liquidación completa")
    _c(
        doc,
        r"""
def liquidar_estadia(tipo, dias, entrada, salida, lavado, tarifas, configuracion):
    if tipo not in tarifas:
        raise ValueError("Tipo de vehículo inválido.")

    dias = parsear_dias(dias)
    lavado = parsear_lavado(lavado)

    if dias > 0:
        base = dias * tarifas[tipo]["tope_dia"]
        minutos = dias * configuracion["minutos_dia_completo"]
        estacionamiento_cobrado = base
        bonificacion = 0
        entrada_registro = ""
        salida_registro = ""
    else:
        minutos = minutos_entre(entrada, salida)
        base = calcular_estacionamiento(minutos, tipo, tarifas, configuracion)
        if lavado and minutos <= configuracion["minutos_limite_primer_tramo"]:
            estacionamiento_cobrado = 0
        elif lavado:
            estacionamiento_cobrado = calcular_excedente(
                minutos, tipo, tarifas, configuracion
            )
        else:
            estacionamiento_cobrado = base
        bonificacion = base - estacionamiento_cobrado
        entrada_registro = entrada
        salida_registro = salida

    valor_lavado = tarifas[tipo]["precio_lavado"] if lavado else 0
    total_cobrado = estacionamiento_cobrado + valor_lavado

    return {
        "dias": dias,
        "minutos": minutos,
        "base": base,
        "bonificacion": bonificacion,
        "lavado": valor_lavado,
        "total_cobrado": total_cobrado,
        "entrada": entrada_registro,
        "salida": salida_registro,
    }
""",
    )

    _h(doc, "8. Casos verificados")
    _t(
        doc,
        ["Caso", "Base", "Bonificación", "Lavado", "Total"],
        [
            ["Auto 25 min", "700", "0", "0", "700"],
            ["Auto 135 min", "4.000", "0", "0", "4.000"],
            ["Auto 210 min", "5.200", "0", "0", "5.200"],
            ["Auto 210 min + lavado", "5.200", "2.800", "5.000", "7.400"],
            ["Moto 45 min + lavado", "800", "800", "2.500", "2.500"],
            ["Camioneta 2 días + lavado", "48.000", "0", "7.500", "55.500"],
        ],
    )

    _h(doc, "9. Cómo defender esta fase")
    _b(
        doc,
        [
            "La recursión no es un adorno: separa naturalmente los dos tramos.",
            "El acumulador minutos_tarifados indica en qué tramo está la función.",
            "El caso base corta la recursión cuando ya no quedan minutos.",
            "El tope se aplica después del cálculo por tiempo.",
            "El excedente con lavado se calcula solo sobre los minutos posteriores a 120.",
            "Para días > 0 no se usa la recursión: se multiplica días por tope.",
        ],
    )

    _h(doc, "10. Preguntas del docente")
    _qa(
        doc,
        [
            (
                "¿Por qué recursión?",
                "Porque el cálculo se repite bloque por bloque y la recursión permite "
                "expresar esa repetición con un caso base claro.",
            ),
            (
                "¿Qué hace minutos_tarifados?",
                "Indica cuántos minutos ya se cobraron para saber si corresponde un "
                "bloque de 30 o una hora adicional.",
            ),
            (
                "¿Cómo se aplica el tope?",
                "Con min(importe_calculado, tope_dia) después de la recursión.",
            ),
            (
                "¿Cómo se calcula el lavado con más de 120 minutos?",
                "Se bonifican los primeros 120 y se cobra el excedente como horas "
                "adicionales, más el precio fijo del lavado.",
            ),
            (
                "¿Qué pasa si los minutos son 0?",
                "Se rechaza antes de calcular porque la salida debe ser posterior a la entrada.",
            ),
            (
                "¿Y si el auto supera el tope?",
                "Se cobra el tope. Por ejemplo, 1000 minutos de auto se cobran $16.000.",
            ),
        ],
    )

    _h(doc, "11. Checklist de la fase")
    _check(
        doc,
        [
            "135 minutos de auto dan $4.000.",
            "210 minutos de auto con lavado dan $7.400.",
            "Moto 45 minutos con lavado da $2.500.",
            "Camioneta 2 días con lavado da $55.500.",
            "El tope de auto se aplica con 1000 minutos.",
            "Días > 0 no piden hora de salida.",
        ],
    )

    _guardar(doc, "03_calculo_recursion.docx")


def documento_04():
    doc = _nuevo(
        "04 de 09",
        "Persistencia en CSV",
        "Reportes del lote y acumulación del cobro manual",
    )

    _h(doc, "1. Objetivo de la fase")
    _p(
        doc,
        "Guardar los resultados en archivos CSV legibles por personas y por planillas. "
        "El lote sobrescribe sus reportes; el cobro manual acumula los registros.",
    )

    _h(doc, "2. Archivos CSV")
    _t(
        doc,
        ["Archivo", "Modo", "Contenido"],
        [
            ["liquidacion.csv", "Sobrescribir", "Movimientos válidos del lote"],
            ["rechazados.csv", "Sobrescribir", "Movimientos inválidos y su motivo"],
            ["liquidacion-manual.csv", "Acumular", "Cobros manuales históricos"],
        ],
    )

    _h(doc, "3. Columnas")
    _t(
        doc,
        ["CSV", "Columnas"],
        [
            [
                "liquidacion.csv",
                "id, fecha_operacion, patente, tipo_vehiculo, entrada, salida, "
                "minutos, base, bonificacion, lavado, total_cobrado",
            ],
            [
                "liquidacion-manual.csv",
                "id, fecha_operacion, patente, tipo_vehiculo, dias, entrada, salida, "
                "minutos, base, bonificacion, lavado, total_cobrado",
            ],
            [
                "rechazados.csv",
                "id, patente, tipo_vehiculo, entrada, salida, motivo",
            ],
        ],
    )

    _h(doc, "4. Sobrescribir vs. acumular")
    _p(
        doc,
        "El lote se procesa completo cada vez, por eso sus CSV se abren en modo w y "
        "se vuelven a generar. El cobro manual, en cambio, representa transacciones "
        "sucesivas: se abre en modo a para no perder las anteriores.",
    )
    _c(
        doc,
        r"""
def agregar_liquidacion_manual(registro, ruta=CSV_MANUAL):
    ruta = Path(ruta)
    existe = ruta.exists() and ruta.stat().st_size > 0
    with open(ruta, "a", newline="", encoding="utf-8") as archivo:
        escritor = csv.DictWriter(
            archivo,
            fieldnames=COLUMNAS_MANUAL,
            delimiter=SEPARADOR,
            extrasaction="ignore",
        )
        if not existe:
            escritor.writeheader()
        escritor.writerow(_formatear_registro(registro, COLUMNAS_MANUAL))
""",
    )

    _h(doc, "5. Encabezado único")
    _p(
        doc,
        "Antes de abrir el archivo se comprueba si existe y si tiene contenido. Si no "
        "existe o está vacío, se escribe el encabezado. Si ya tiene filas, se agrega "
        "solo el registro nuevo.",
    )

    _h(doc, "6. Formato de montos")
    _p(
        doc,
        "En pantalla se muestra $ 7.400,00. En el CSV se guarda 7400.00 para que el "
        "valor sea numérico y fácil de sumar. La GUI y la consola usan formato_pesos; "
        "csv_repo usa dos decimales sin símbolo.",
    )
    _c(
        doc,
        r"""
def _formatear_registro(registro, columnas):
    fila = {}
    for columna in columnas:
        valor = registro.get(columna, "")
        if columna in CAMPOS_MONTO and valor != "":
            fila[columna] = f"{float(valor):.2f}"
        else:
            fila[columna] = valor
    return fila
""",
    )

    _h(doc, "7. Ejemplo de fila manual")
    _c(
        doc,
        r"""
MAN-20261025183000123456,2026-10-25,AF567GG,auto,0,09:00,12:30,
210,5200.00,2800.00,5000.00,7400.00
""",
    )
    _nota(
        doc,
        "La fila real va en una sola línea. Aquí se muestra partida para que se lea "
        "en el documento.",
    )

    _h(doc, "8. Manejo de errores")
    _p(
        doc,
        "Si el archivo está abierto en Excel, Windows puede bloquear la escritura. "
        "El sistema captura OSError y muestra un mensaje para que el usuario cierre "
        "el archivo y vuelva a intentar. No se pierde el cálculo ya realizado.",
    )

    _h(doc, "9. Cómo defender esta fase")
    _b(
        doc,
        [
            "El lote genera reportes completos y reproducibles.",
            "El manual acumula porque son operaciones independientes.",
            "El encabezado se escribe una sola vez.",
            "Los montos quedan numéricos en el CSV.",
            "El separador está centralizado en config.py.",
            "Los errores de escritura se informan, no se ocultan.",
        ],
    )

    _h(doc, "10. Preguntas del docente")
    _qa(
        doc,
        [
            (
                "¿Por qué el manual es append?",
                "Porque cada cobro es una transacción nueva y no debe borrar las anteriores.",
            ),
            (
                "¿Cómo evitás duplicar el encabezado?",
                "Se escribe solo si el archivo no existe o está vacío.",
            ),
            (
                "¿Qué pasa si el archivo no existe?",
                "Se crea automáticamente con el encabezado.",
            ),
            (
                "¿Qué pasa si está abierto en Excel?",
                "Se captura OSError y se muestra un mensaje para cerrarlo.",
            ),
            (
                "¿Por qué coma?",
                "Es el separador estándar y está centralizado por si se necesita cambiar.",
            ),
        ],
    )

    _h(doc, "11. Checklist de la fase")
    _check(
        doc,
        [
            "liquidacion.csv tiene 30 filas después del lote.",
            "rechazados.csv tiene 3 filas.",
            "liquidacion-manual.csv acumula dos registros seguidos.",
            "El encabezado aparece una sola vez.",
            "Los montos tienen dos decimales.",
        ],
    )

    _guardar(doc, "04_csv.docx")


def documento_05():
    doc = _nuevo(
        "05 de 09",
        "Procesamiento por lote",
        "Validación de los 33 movimientos y generación de reportes",
    )

    _h(doc, "1. Objetivo de la fase")
    _p(
        doc,
        "Leer movimientos.json, validar cada movimiento, liquidar los válidos y separar "
        "los inválidos con un motivo. El resultado se guarda en liquidacion.csv y "
        "rechazados.csv.",
    )

    _h(doc, "2. Flujo")
    _n(
        doc,
        [
            "main carga tarifas.json y movimientos.json.",
            "lote.procesar_lote recorre la lista en orden.",
            "Valida ID, patente, tipo, lavado y horario.",
            "Si hay error, agrega el movimiento a rechazados.",
            "Si es válido, llama a calculo.liquidar_estadia.",
            "Acumula totales por tipo y cantidad de lavados.",
            "main escribe los CSV y muestra el resumen.",
        ],
    )

    _h(doc, "3. Validaciones del lote")
    _t(
        doc,
        ["Validación", "Motivo si falla"],
        [
            ["ID presente y no repetido", "ID inválido / ID repetido"],
            ["Patente válida", "Patente inválida"],
            ["Tipo existente en tarifas", "Tipo de vehículo inválido"],
            ["Lavado interpretable", "Valor de lavado inválido"],
            ["Horas válidas y salida posterior", "Hora inválida / salida anterior"],
        ],
    )

    _h(doc, "4. Código principal")
    _c(
        doc,
        r"""
def procesar_lote(movimientos, tarifas, configuracion, fecha_operacion):
    validos = []
    rechazados = []
    ids_vistos = set()

    for movimiento in movimientos:
        id_movimiento = movimiento.get("id")
        patente_original = movimiento.get("patente")
        tipo_original = movimiento.get("tipo_vehiculo")
        entrada = movimiento.get("entrada")
        salida = movimiento.get("salida")
        lavado_raw = movimiento.get("solicito_lavado", False)
        motivo = None
        lavado = False

        if not isinstance(id_movimiento, str) or not id_movimiento.strip():
            motivo = "ID inválido"
        elif id_movimiento in ids_vistos:
            motivo = "ID repetido"
        else:
            ids_vistos.add(id_movimiento)

        patente = normalizar_patente(patente_original)
        tipo = normalizar_tipo(tipo_original, tarifas)

        if motivo is None and not patente_valida(patente):
            motivo = "Patente inválida"
        if motivo is None and tipo is None:
            motivo = "Tipo de vehículo inválido"

        if motivo is not None:
            rechazados.append({...})
            continue

        resultado = liquidar_estadia(tipo, 0, entrada, salida, lavado, tarifas, configuracion)
        validos.append({...})

    return validos, rechazados, resumen
""",
    )
    _nota(doc, "El código completo está en lote.py. Aquí se muestra la estructura.") 

    _h(doc, "5. Resultado esperado")
    _t(
        doc,
        ["Métrica", "Valor"],
        [
            ["Movimientos leídos", "33"],
            ["Procesados", "30"],
            ["Rechazados", "3"],
            ["Total recaudado", "$129.000,00"],
            ["Lavados", "13"],
        ],
    )

    _h(doc, "6. Totales por tipo")
    _t(
        doc,
        ["Tipo", "Cantidad", "Total"],
        [
            ["Auto", "13", "$49.500,00"],
            ["Moto", "8", "$14.600,00"],
            ["Camioneta", "9", "$64.900,00"],
        ],
    )

    _h(doc, "7. Rechazos")
    _t(
        doc,
        ["ID", "Patente", "Motivo"],
        [
            ["MOV031", "BH123JJ", "La salida debe ser posterior a la entrada."],
            ["MOV032", "BJ234KK", "Tipo de vehículo inválido"],
            ["MOV033", "PATENTE99", "Patente inválida"],
        ],
    )

    _h(doc, "8. Cómo defender esta fase")
    _b(
        doc,
        [
            "Se procesa en orden para que el resultado sea reproducible.",
            "Los rechazados no se liquidan ni se escriben en liquidacion.csv.",
            "Cada rechazo conserva su motivo para auditoría.",
            "Las patentes y tipos se normalizan antes de validar.",
            "Se calcula el total recaudado y también totales por tipo.",
            "El lote usa la misma función de cálculo que el cobro manual.",
        ],
    )

    _h(doc, "9. Preguntas del docente")
    _qa(
        doc,
        [
            (
                "¿Por qué se rechaza MOV031?",
                "Porque entra a las 18:00 y sale a las 15:00; la salida es anterior.",
            ),
            (
                "¿Por qué se rechaza MOV032?",
                "Porque 'tractor' no existe en tarifas.json.",
            ),
            (
                "¿Por qué se rechaza MOV033?",
                "Porque PATENTE99 no cumple ninguno de los dos formatos.",
            ),
            (
                "¿Qué pasa con las patentes con espacios?",
                "Se normalizan. Por ejemplo, ' BF 901 GG ' se procesa como BF901GG.",
            ),
            (
                "¿Qué pasa con los tipos en mayúscula?",
                "Se normalizan. 'AUTO' y ' MOTO ' se aceptan.",
            ),
            (
                "¿Por qué hay 13 lavados?",
                "Porque 13 movimientos válidos tienen solicito_lavado en true.",
            ),
        ],
    )

    _h(doc, "10. Checklist de la fase")
    _check(
        doc,
        [
            "Ejecutar la opción 2.",
            "Verificar 33 leídos, 30 procesados y 3 rechazados.",
            "Verificar total $129.000.",
            "Abrir liquidacion.csv y contar 30 filas.",
            "Abrir rechazados.csv y verificar los 3 motivos.",
        ],
    )

    _guardar(doc, "05_lote.docx")


def documento_06():
    doc = _nuevo(
        "06 de 09",
        "Cobro manual por consola y consulta",
        "Respaldo del cobro manual y lectura del historial",
    )

    _h(doc, "1. Objetivo de la fase")
    _p(
        doc,
        "Permitir el cobro manual sin depender de tkinter y ofrecer una consulta "
        "simple de las liquidaciones guardadas. Es el respaldo si la GUI no abre en "
        "la máquina de la defensa.",
    )

    _h(doc, "2. Flujo del registro")
    _n(
        doc,
        [
            "Pedir patente hasta que sea válida.",
            "Pedir tipo hasta que exista en tarifas.",
            "Preguntar días.",
            "Si días = 0, pedir entrada y salida.",
            "Si días > 0, informar que se liquida por días completos.",
            "Preguntar lavado.",
            "Calcular, mostrar el desglose y guardar en liquidacion-manual.csv.",
        ],
    )

    _h(doc, "3. Bucles de validación")
    _p(
        doc,
        "Cada dato se pide dentro de un while. Si la validación falla, se muestra el "
        "error y se vuelve a preguntar. Así el usuario no puede avanzar con datos "
        "inválidos.",
    )
    _c(
        doc,
        r"""
def pedir_patente():
    while True:
        patente = normalizar_patente(input("Patente (ABC123 o AA123BB): "))
        if patente_valida(patente):
            return patente
        print("Patente inválida. Usá ABC123 o AA123BB.")


def pedir_tipo(tarifas):
    while True:
        tipo = normalizar_tipo(input("Tipo (auto/moto/camioneta): "), tarifas)
        if tipo is not None:
            return tipo
        print("Tipo de vehículo inválido.")
""",
    )

    _h(doc, "4. Horas y días")
    _p(
        doc,
        "Las horas se normalizan a HH:MM. Si los días son mayores a 0, no se piden "
        "horarios porque se cobra por jornada completa.",
    )
    _c(
        doc,
        r"""
if dias > 0:
    print("Se liquidará por días completos a costo fijo.")
    entrada = ""
    salida = ""
else:
    entrada = pedir_hora("Hora de entrada (HH:MM): ")
    salida = pedir_hora("Hora de salida (HH:MM): ")
    minutos_entre(entrada, salida)
""",
    )

    _h(doc, "5. Desglose en pantalla")
    _p(
        doc,
        "Antes de guardar se muestra patente, tipo, días o minutos, base, bonificación, "
        "lavado y total. Para 210 minutos de auto con lavado, el desglose es:",
    )
    _t(
        doc,
        ["Concepto", "Valor"],
        [
            ["Minutos", "210"],
            ["Base", "$5.200,00"],
            ["Bonificación", "$2.800,00"],
            ["Lavado", "$5.000,00"],
            ["Total", "$7.400,00"],
        ],
    )

    _h(doc, "6. Consulta del historial")
    _p(
        doc,
        "La consulta lee liquidacion-manual.csv y muestra las últimas diez filas. Si el "
        "archivo no existe o está vacío, informa que todavía no hay registros.",
    )
    _c(
        doc,
        r"""
def consultar_liquidaciones():
    if not CSV_MANUAL.exists() or CSV_MANUAL.stat().st_size == 0:
        print("No hay liquidaciones manuales registradas.")
        return
    with open(CSV_MANUAL, "r", newline="", encoding="utf-8") as archivo:
        filas = list(csv.DictReader(archivo))
    for fila in filas[-10:]:
        print(fila["id"], fila["patente"], fila["tipo_vehiculo"], fila["total_cobrado"])
""",
    )

    _h(doc, "7. Cómo defender esta fase")
    _b(
        doc,
        [
            "La consola es el respaldo si tkinter no funciona.",
            "Los bucles evitan avanzar con datos inválidos.",
            "La consulta demuestra que el CSV se acumula.",
            "La lógica de cálculo es la misma que usa la GUI.",
            "No se duplica el cálculo ni la validación.",
        ],
    )

    _h(doc, "8. Preguntas del docente")
    _qa(
        doc,
        [
            (
                "¿Cómo evitás bucles infinitos?",
                "El usuario siempre puede escribir un dato válido; cada error muestra "
                "el formato esperado. Además, Ctrl+C corta el programa de forma limpia.",
            ),
            (
                "¿Qué pasa si días > 0?",
                "No se piden horas y se cobra días * tope_dia.",
            ),
            (
                "¿Dónde se guarda?",
                "En liquidacion-manual.csv, en modo append.",
            ),
            (
                "¿Qué pasa si no hay CSV?",
                "La consulta informa que no hay registros.",
            ),
            (
                "¿Por qué mantener la consola si existe la GUI?",
                "Porque es un respaldo y porque la consigna base pide una aplicación de consola.",
            ),
        ],
    )

    _h(doc, "9. Checklist de la fase")
    _check(
        doc,
        [
            "Registrar un cobro por consola con 210 minutos y lavado.",
            "Verificar que el total sea $7.400.",
            "Consultar el historial y ver la fila.",
            "Probar una patente inválida y un tipo inválido.",
            "Probar días = 2 y confirmar que no pide horas.",
        ],
    )

    _guardar(doc, "06_manual_consola_consulta.docx")


def documento_07():
    doc = _nuevo(
        "07 de 09",
        "Interfaz gráfica con tkinter",
        "Desafío de nivel 2: ventana, eventos y total en tiempo real",
    )

    _h(doc, "1. Objetivo de la fase")
    _p(
        doc,
        "Construir una ventana de escritorio para el Cobro Manual en Ventanilla. La "
        "ventana captura los datos, muestra el total mientras se escribe y guarda la "
        "operación en liquidacion-manual.csv.",
    )
    _nota(
        doc,
        "Este es el desafío de nivel 2. No reemplaza la lógica: la reutiliza. "
        "La GUI es una capa de presentación sobre validaciones, cálculo y CSV.",
    )

    _h(doc, "2. Por qué tkinter")
    _b(
        doc,
        [
            "Es nativo de Python: no requiere instalar paquetes externos.",
            "Incluye Entry, Combobox, Checkbutton, Spinbox y messagebox.",
            "Permite demostrar programación orientada a eventos.",
            "Es suficiente para una ventana de ventanilla.",
        ],
    )

    _h(doc, "3. Layout de la ventana")
    _t(
        doc,
        ["Zona", "Componente", "Función"],
        [
            ["Datos", "Entry patente", "Capturar la patente"],
            ["Datos", "Combobox tipo", "Elegir auto, moto o camioneta"],
            ["Datos", "Spinbox días", "Elegir 0 o más días"],
            ["Datos", "Label aviso", "Informar si se cobra por días completos"],
            ["Datos", "Entry entrada/salida", "Capturar HH:MM"],
            ["Datos", "Checkbutton lavado", "Indicar si solicitó lavado"],
            ["Resultado", "Labels", "Mostrar base, bonificación, lavado y total"],
            ["Acciones", "Botones", "Liquidar y Guardar, Limpiar, Cerrar"],
        ],
    )

    _h(doc, "4. Variables de control")
    _p(
        doc,
        "Cada campo se conecta a una StringVar o BooleanVar. Cuando la variable cambia, "
        "tkinter puede avisar y actualizar el total sin que el usuario presione un botón.",
    )
    _c(
        doc,
        r"""
self.var_patente = tk.StringVar()
self.var_tipo = tk.StringVar()
self.var_dias = tk.StringVar(value="0")
self.var_entrada = tk.StringVar()
self.var_salida = tk.StringVar()
self.var_lavado = tk.BooleanVar(value=False)
self.var_aviso = tk.StringVar(value="")
self.var_total = tk.StringVar(value="—")
""",
    )

    _h(doc, "5. Eventos")
    _p(
        doc,
        "La ventana usa mainloop. Los eventos se conectan con trace_add, bind y command. "
        "Cada cambio dispara actualizar_total.",
    )
    _c(
        doc,
        r"""
def conectar_eventos(self):
    self.var_patente.trace_add("write", self.actualizar_total)
    self.var_tipo.trace_add("write", self.actualizar_total)
    self.var_dias.trace_add("write", self._al_cambiar_dias)
    self.var_entrada.trace_add("write", self.actualizar_total)
    self.var_salida.trace_add("write", self.actualizar_total)
    self.entry_patente.bind("<Return>", lambda event: self.liquidar())
""",
    )

    _h(doc, "6. Total en tiempo real")
    _p(
        doc,
        "actualizar_total intenta leer y liquidar. Si falta un dato o es inválido, "
        "muestra guiones y no interrumpe con un messagebox. El error recién se muestra "
        "al presionar Liquidar y Guardar.",
    )
    _c(
        doc,
        r"""
def actualizar_total(self, *args):
    if self.actualizando:
        return
    self.actualizando = True
    try:
        datos = self._leer_formulario()
        resultado = liquidar_estadia(
            datos["tipo"], datos["dias"], datos["entrada"], datos["salida"],
            datos["lavado"], self.tarifas, self.configuracion,
        )
    except (ValueError, TypeError):
        self._mostrar_guiones()
        return
    finally:
        self.actualizando = False

    self.var_minutos.set(f"{resultado['minutos']} min")
    self.var_base.set(formato_pesos(resultado["base"]))
    self.var_bonificacion.set(formato_pesos(resultado["bonificacion"]))
    self.var_lavado_monto.set(formato_pesos(resultado["lavado"]))
    self.var_total.set(formato_pesos(resultado["total_cobrado"]))
""",
    )

    _h(doc, "7. Días mayores a 0")
    _p(
        doc,
        "Cuando el usuario pone días > 0, la ventana deshabilita entrada y salida, "
        "muestra el tiempo en días y muestra un aviso explícito de que se cobra por "
        "días completos. El cálculo usa días * tope_dia, sin horas.",
    )
    _c(
        doc,
        r"""
def actualizar_estado_horas(self):
    try:
        dias = parsear_dias(self.var_dias.get())
    except ValueError:
        dias = 0
    estado = "disabled" if dias > 0 else "normal"
    self.entry_entrada.configure(state=estado)
    self.entry_salida.configure(state=estado)
    if dias > 0:
        self.var_aviso.set("Se liquida por días completos a costo fijo.")
    else:
        self.var_aviso.set("")
""",
    )

    _h(doc, "8. Guardar y mostrar el cartel")
    _p(
        doc,
        "liquidar() valida, calcula, guarda en el CSV y muestra un messagebox con el "
        "desglose. Si el archivo está bloqueado, muestra un error y no cierra la ventana.",
    )
    _c(
        doc,
        r"""
def liquidar(self):
    try:
        datos = self._leer_formulario()
        resultado = liquidar_estadia(
            datos["tipo"], datos["dias"], datos["entrada"], datos["salida"],
            datos["lavado"], self.tarifas, self.configuracion,
        )
    except (ValueError, TypeError) as error:
        messagebox.showerror("Datos inválidos", str(error), parent=self.root)
        return

    registro = {...}
    try:
        agregar_liquidacion_manual(registro)
    except OSError as error:
        messagebox.showerror("Error al guardar", str(error), parent=self.root)
        return

    messagebox.showinfo("Liquidación registrada", mensaje, parent=self.root)
    self.limpiar()
""",
    )

    _h(doc, "9. Respaldo si tkinter falla")
    _p(
        doc,
        "main.py importa gui_manual dentro de la opción 1. Si tkinter no está "
        "disponible, captura ImportError y sugiere usar la opción 2, que es el cobro "
        "por consola. El programa no se cae.",
    )

    _h(doc, "10. Cómo defender esta fase")
    _b(
        doc,
        [
            "La GUI es una capa de presentación, no una copia del sistema.",
            "Los eventos se conectan a variables y a funciones concretas.",
            "El total se actualiza sin presionar el botón.",
            "El botón valida, guarda y muestra el cartel flotante.",
            "Los días > 0 deshabilitan los horarios.",
            "Si el CSV está bloqueado, se informa el error.",
            "La consola sigue disponible como respaldo.",
        ],
    )

    _h(doc, "11. Preguntas del docente")
    _qa(
        doc,
        [
            (
                "¿Qué es mainloop?",
                "El ciclo de eventos de tkinter: espera clics, teclas y cambios y "
                "despacha las funciones conectadas.",
            ),
            (
                "¿Qué es una StringVar?",
                "Una variable de tkinter que guarda texto y se puede conectar a un widget.",
            ),
            (
                "¿Cómo se actualiza el total?",
                "Con trace_add sobre las variables, bind para Enter y command en el checkbox.",
            ),
            (
                "¿Dónde está la lógica de cálculo?",
                "En calculo.py. La GUI solo llama a liquidar_estadia.",
            ),
            (
                "¿Cómo evitás duplicar la validación?",
                "La GUI usa las mismas funciones de validaciones.py.",
            ),
            (
                "¿Qué pasa si el usuario cierra la ventana?",
                "Se destruye la ventana y main vuelve al submenú.",
            ),
            (
                "¿Qué pasa si tkinter no está instalado?",
                "main captura ImportError y ofrece el cobro por consola.",
            ),
            (
                "¿Cómo se guarda?",
                "Con csv_repo.agregar_liquidacion_manual, en modo append.",
            ),
        ],
    )

    _h(doc, "12. Checklist de la fase")
    _check(
        doc,
        [
            "Abrir la ventana desde el submenú.",
            "Cargar auto con 210 minutos y lavado; ver $7.400.",
            "Guardar y ver el cartel flotante.",
            "Abrir liquidacion-manual.csv y ver la fila.",
            "Cargar 2 días de camioneta con lavado; ver $55.500.",
            "Probar una patente inválida.",
            "Cerrar la ventana y volver al menú.",
        ],
    )

    _guardar(doc, "07_gui_tkinter.docx")


def documento_08():
    doc = _nuevo(
        "08 de 09",
        "Menú e integración",
        "Unión de módulos, submenú manual y manejo de errores",
    )

    _h(doc, "1. Objetivo de la fase")
    _p(
        doc,
        "Unir todos los módulos desde main.py, respetar el menú de la consigna y "
        "ofrecer un submenú para elegir ventana gráfica o consola.",
    )

    _h(doc, "2. Menú principal")
    _c(
        doc,
        r"""
==================================================
SISTEMA DE GESTIÓN DE ESTACIONAMIENTO AUTOPARK
==================================================
1. Registrar / Consultar cobro manual
2. Procesar lote de movimientos (JSON -> CSV)
3. Salir
Seleccione una opción:
""",
    )

    _h(doc, "3. Submenú de cobro manual")
    _c(
        doc,
        r"""
1. Registrar cobro con ventana gráfica
2. Registrar cobro por consola
3. Consultar liquidaciones manuales
4. Volver
""",
    )
    _p(
        doc,
        "Se eligió un submenú para cumplir 'Registrar / Consultar' y, al mismo tiempo, "
        "mantener la consola como respaldo.",
    )

    _h(doc, "4. Flujo de cada opción")
    _t(
        doc,
        ["Opción", "Acción"],
        [
            ["1", "Abre el submenú de cobro manual"],
            ["1.1", "Importa gui_manual y abre la ventana"],
            ["1.2", "Ejecuta el cobro por consola"],
            ["1.3", "Lee y muestra liquidacion-manual.csv"],
            ["2", "Carga movimientos.json, procesa y escribe CSV"],
            ["3", "Termina el programa"],
        ],
    )

    _h(doc, "5. Carga de tarifas")
    _p(
        doc,
        "main carga tarifas.json una sola vez y pasa configuración y tarifas a la "
        "consola, la GUI y el lote. Si el archivo falta o es inválido, el programa "
        "informa y termina, porque no puede calcular sin precios.",
    )
    _c(
        doc,
        r"""
def main():
    try:
        configuracion, tarifas = cargar_tarifas()
    except (OSError, ValueError) as error:
        print(f"Error al cargar tarifas.json: {error}")
        return 1
    while True:
        opcion = input("Seleccione una opción: ").strip()
        if opcion == "1":
            menu_manual(tarifas, configuracion)
        elif opcion == "2":
            procesar_lote_menu(tarifas, configuracion)
        elif opcion == "3":
            return 0
        else:
            print("Opción inválida. Ingrese 1, 2 o 3.")
""",
    )

    _h(doc, "6. Importación segura de tkinter")
    _p(
        doc,
        "gui_manual se importa dentro de la opción 1.1, no al inicio de main. Si tkinter "
        "no está disponible, el error se captura y el usuario puede seguir por consola.",
    )
    _c(
        doc,
        r"""
try:
    from gui_manual import abrir_ventana_manual
except ImportError:
    print("GUI no disponible. Usá la opción 2.")
    continue
try:
    abrir_ventana_manual(tarifas, configuracion)
except Exception as error:
    print(f"No se pudo abrir la ventana: {error}")
""",
    )

    _h(doc, "7. Manejo de errores")
    _b(
        doc,
        [
            "Opción inválida: se informa y se vuelve a pedir.",
            "tarifas.json faltante o inválido: mensaje y salida.",
            "movimientos.json faltante o inválido: mensaje y vuelta al menú.",
            "CSV bloqueado: mensaje y vuelta al menú.",
            "Ctrl+C o EOF: salida limpia sin traceback.",
            "GUI no disponible: se sugiere la consola.",
        ],
    )

    _h(doc, "8. Código de cierre limpio")
    _c(
        doc,
        r"""
if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (KeyboardInterrupt, EOFError):
        print("\nSaliendo...")
        raise SystemExit(0)
""",
    )

    _h(doc, "9. Cómo defender esta fase")
    _b(
        doc,
        [
            "main es el punto de entrada y no contiene reglas de negocio.",
            "El menú respeta las tres opciones de la consigna.",
            "El submenú agrega la GUI sin romper la consola.",
            "Las tarifas se cargan una vez y se comparten.",
            "La GUI se importa de forma segura.",
            "Los errores se informan sin cerrar abruptamente.",
        ],
    )

    _h(doc, "10. Preguntas del docente")
    _qa(
        doc,
        [
            (
                "¿Por qué el menú tiene tres opciones?",
                "Porque es el menú pedido en el enunciado. El submenú agrega la GUI "
                "dentro de la opción 1.",
            ),
            (
                "¿Dónde se abre la GUI?",
                "En la opción 1.1 del submenú de cobro manual.",
            ),
            (
                "¿Qué pasa si no hay tkinter?",
                "Se captura ImportError y se sugiere la opción 1.2, cobro por consola.",
            ),
            (
                "¿Qué pasa si falta movimientos.json?",
                "Se informa el error y se vuelve al menú; el programa no se cierra.",
            ),
            (
                "¿Cómo se vuelve al menú después de cerrar la GUI?",
                "mainloop termina al destruir la ventana y la función retorna al submenú.",
            ),
        ],
    )

    _h(doc, "11. Checklist de la fase")
    _check(
        doc,
        [
            "Ejecutar main.py y ver el menú.",
            "Entrar a la opción 1 y ver el submenú.",
            "Abrir la GUI y cerrarla; volver al submenú.",
            "Ejecutar la opción 2 y volver al menú.",
            "Probar una opción inválida.",
            "Probar Ctrl+C.",
        ],
    )

    _guardar(doc, "08_menu_integracion.docx")


def documento_09():
    doc = _nuevo(
        "09 de 09",
        "Pruebas y defensa final",
        "Casos de prueba, guion de demo y preguntas frecuentes",
    )

    _h(doc, "1. Objetivo de la fase")
    _p(
        doc,
        "Comprobar que el sistema funciona antes de la defensa y tener un recorrido "
        "claro para mostrar. Las pruebas automáticas cubren la lógica; las pruebas "
        "manuales cubren la GUI.",
    )

    _h(doc, "2. Estrategia de pruebas")
    _b(
        doc,
        [
            "Pruebas automáticas con assert para validaciones y cálculo.",
            "Prueba de integración del lote completo.",
            "Prueba de escritura y acumulación de CSV en carpeta temporal.",
            "Prueba manual de la ventana y de los mensajes.",
        ],
    )

    _h(doc, "3. Casos automáticos")
    _t(
        doc,
        ["Caso", "Resultado esperado"],
        [
            ["' BF 901 GG '", "BF901GG"],
            ["AA123BB", "Patente válida"],
            ["PATENTE99", "Patente inválida"],
            ["' MOTO '", "moto"],
            ["8:00", "480 minutos"],
            ["135 min auto", "$4.000"],
            ["210 min auto + lavado", "$7.400"],
            ["Moto 45 min + lavado", "$2.500"],
            ["Camioneta 2 días + lavado", "$55.500"],
            ["Lote completo", "30 válidos, 3 rechazados, $129.000"],
        ],
    )

    _h(doc, "4. Resultado del lote")
    _t(
        doc,
        ["Métrica", "Valor"],
        [
            ["Válidos", "30"],
            ["Rechazados", "3"],
            ["Total", "$129.000,00"],
            ["Auto", "13 / $49.500,00"],
            ["Moto", "8 / $14.600,00"],
            ["Camioneta", "9 / $64.900,00"],
            ["Lavados", "13"],
        ],
    )

    _h(doc, "5. Pruebas manuales de la GUI")
    _n(
        doc,
        [
            "Abrir main.py y entrar a la opción 1.1.",
            "Verificar que la ventana abra y que el foco esté en patente.",
            "Cargar AA123BB, auto, 0 días, 09:00 a 12:30, lavado.",
            "Verificar base $5.200, bonificación $2.800, lavado $5.000 y total $7.400.",
            "Presionar Liquidar y Guardar y ver el cartel flotante.",
            "Abrir liquidacion-manual.csv y verificar la fila.",
            "Cargar camioneta, 2 días, lavado; verificar que no pida horas y total $55.500.",
            "Probar PATENTE99 y ver el error.",
            "Cerrar la ventana y verificar que vuelve al submenú.",
        ],
    )

    _h(doc, "6. Guion de demo")
    _t(
        doc,
        ["Paso", "Qué mostrar", "Qué explicar"],
        [
            ["1", "Estructura del proyecto", "Separación de responsabilidades"],
            ["2", "Opción 2", "30 válidos, 3 rechazados, $129.000"],
            ["3", "liquidacion.csv y rechazados.csv", "Persistencia del lote"],
            ["4", "GUI con 210 min + lavado", "Total en vivo y cálculo"],
            ["5", "Cartel y CSV manual", "Guardado acumulativo"],
            ["6", "Camioneta 2 días", "Días > 0 sin horas"],
            ["7", "Patente inválida", "Validación y mensaje"],
        ],
    )

    _h(doc, "7. Preguntas frecuentes de la defensa")
    _qa(
        doc,
        [
            (
                "¿Qué parte es el desafío de nivel 2?",
                "La ventana tkinter para el cobro manual, con widgets, eventos y total "
                "en tiempo real.",
            ),
            (
                "¿La GUI duplica el cálculo?",
                "No. Llama a calculo.liquidar_estadia y a csv_repo.agregar_liquidacion_manual.",
            ),
            (
                "¿Por qué no hay SQLite?",
                "Porque el desafío elegido es GUI y el CSV cumple la persistencia base.",
            ),
            (
                "¿Cómo se calculan los 135 minutos?",
                "4 bloques de 30 a $700 y 1 hora adicional a $1.200: $4.000.",
            ),
            (
                "¿Cómo se aplica el tope?",
                "Con min(importe, tope_dia) después de la recursión.",
            ),
            (
                "¿Qué pasa con el lavado y 210 minutos?",
                "Se bonifican los primeros 120 y se cobran 2 horas excedentes: $2.400 "
                "más $5.000 de lavado = $7.400.",
            ),
            (
                "¿Cómo se guarda el manual?",
                "En liquidacion-manual.csv en modo append, con encabezado solo la primera vez.",
            ),
            (
                "¿Qué pasa si el CSV está abierto en Excel?",
                "Se captura OSError y se muestra un mensaje para cerrarlo.",
            ),
            (
                "¿Qué pasa si tkinter no abre?",
                "Se captura el error y se usa el cobro por consola.",
            ),
            (
                "¿Cómo se prueba el sistema?",
                "python pruebas.py y la lista de pruebas manuales de la GUI.",
            ),
        ],
    )

    _h(doc, "8. Problemas comunes y solución")
    _t(
        doc,
        ["Problema", "Solución"],
        [
            ["No encuentra tarifas.json", "Ejecutar desde autopark o revisar BASE_DIR"],
            ["tkinter no abre", "Usar la opción 1.2 de consola"],
            ["No puede escribir el CSV", "Cerrar el archivo en Excel"],
            ["El total no se actualiza", "Revisar trace_add y que el dato sea válido"],
            ["El lote da otro total", "Revisar normalización y recursión"],
            [
                "Acentos raros en consola",
                "main configura el code page 65001 y stdout en UTF-8",
            ],
        ],
    )

    _h(doc, "9. Checklist final")
    _check(
        doc,
        [
            "python pruebas.py imprime 'Todas las pruebas pasaron.'",
            "La opción 2 da 30/3/$129.000.",
            "La GUI abre y calcula en vivo.",
            "El manual guarda y acumula.",
            "La consulta muestra los registros.",
            "Los días > 0 no piden horas.",
            "La patente inválida muestra error.",
            "El proyecto no depende de librerías externas.",
            "Los 10 documentos están en documentacion/.",
        ],
    )

    _h(doc, "10. Frases para la defensa")
    _b(
        doc,
        [
            "“La GUI es una capa de presentación; la lógica está en calculo.py.”",
            "“Normalizamos antes de validar para no rechazar datos correctos.”",
            "“El acumulador indica en qué tramo estamos: 30 minutos u hora.”",
            "“El lote y el manual comparten la misma función de liquidación.”",
            "“El CSV manual se acumula porque cada cobro es una transacción.”",
            "“Si tkinter falla, la consola sigue funcionando.”",
        ],
    )

    _guardar(doc, "09_pruebas_defensa.docx")


def main():
    documento_00()
    documento_01()
    documento_02()
    documento_03()
    documento_04()
    documento_05()
    documento_06()
    documento_07()
    documento_08()
    documento_09()
    print("Documentación generada correctamente.")


if __name__ == "__main__":
    main()
